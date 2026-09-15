from __future__ import annotations

import re
import unicodedata
from datetime import datetime
from typing import Any

import pandas as pd
from openpyxl.cell.cell import Cell
from openpyxl.utils.datetime import from_excel

from ..config import ParserConfig
from ..ingestion.excel_reader import ExcelSource
from ..models import ParseResult, ValidationIssue
from .hierarchy import ENTITY_COLUMNS, EntityNode, EntityTreeBuilder


DATE_PATTERN = re.compile(r"(?<!\d)(\d{1,2})[/-](\d{1,2})[/-](\d{4})(?!\d)")
OUTPUT_COLUMNS = [
    "source_file", "source_hash", "sheet_name", "cell_address", "source_row",
    "project_id", "project_label", "entity_id", "parent_entity_id", "entity_level",
    "entity_depth", "entity_label", "entity_path", "entity_key",
    "unit_raw", "unit_original", "unit_normalized", "effective_unit",
    "unit_source_level", "unit_source_entity_id", "parser_rule", "parser_confidence",
    "project", "section", "item", "unit", "date", "metric_original", "metric_normalized", "raw_value",
    "value_numeric", "chart_value", "display_value", "number_format",
    "value_kind", "validation_status",
]


def _ancestor_label(node: EntityNode, level: str, nodes: dict[str, EntityNode]) -> str | None:
    current: EntityNode | None = node
    while current is not None:
        if current.entity_level == level:
            return current.entity_label
        current = nodes.get(current.parent_entity_id) if current.parent_entity_id else None
    return None


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\u00a0", " ").split()).strip()


def fold_text(value: Any) -> str:
    text = unicodedata.normalize("NFD", clean_text(value).casefold())
    return "".join(char for char in text if unicodedata.category(char) != "Mn")


def extract_date(value: Any, epoch: datetime) -> pd.Timestamp | None:
    if isinstance(value, (datetime, pd.Timestamp)):
        return pd.Timestamp(value).normalize()
    if isinstance(value, (int, float)) and 1 <= value <= 2958465:
        return pd.Timestamp(from_excel(value, epoch=epoch)).normalize()
    match = DATE_PATTERN.search(clean_text(value))
    if not match:
        return None
    day, month, year = map(int, match.groups())
    try:
        return pd.Timestamp(year=year, month=month, day=day)
    except ValueError:
        return None


def normalize_metric(metric: str, config: ParserConfig) -> str:
    cleaned = clean_text(metric)
    alias = config.metric_aliases.get(cleaned.casefold())
    if alias:
        return alias
    folded = fold_text(cleaned)
    if cleaned == "%" or "%" in cleaned or "ty le" in folded:
        return "% báo sai"
    if folded.startswith("tong so"):
        return "Tổng số"
    if "bao sai" in folded or "nghi ngo gian lan" in folded:
        return "Báo sai/Lỗi"
    if "ghi chu" in folded:
        return "Ghi chú"
    return cleaned


def _is_blank(value: Any, config: ParserConfig) -> bool:
    if value is None:
        return True
    return clean_text(value).casefold() in {clean_text(v).casefold() for v in config.blank_markers}


def _display_value(cell: Cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool) and "%" in cell.number_format:
        decimal_match = re.search(r"\.([0#]+)%", cell.number_format)
        decimals = len(decimal_match.group(1)) if decimal_match else 0
        return f"{value * 100:.{decimals}f}%"
    return str(value)


def _value_fields(cell: Cell, metric: str, config: ParserConfig) -> dict[str, Any]:
    value = cell.value
    cleaned = clean_text(value)
    is_percent = "%" in cell.number_format or metric == "% báo sai"
    if cleaned.casefold() in {marker.casefold() for marker in config.missing_markers}:
        return {
            "value_numeric": None, "chart_value": None, "display_value": cleaned,
            "value_kind": "missing_marker", "validation_status": "warning",
        }
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        numeric = float(value)
        chart_value = numeric * 100 if "%" in cell.number_format else numeric
        return {
            "value_numeric": numeric, "chart_value": chart_value,
            "display_value": _display_value(cell), "value_kind": "percentage" if is_percent else "numeric",
            "validation_status": "valid",
        }
    if isinstance(value, str) and value.strip().endswith("%"):
        try:
            parsed = float(value.strip()[:-1].replace(",", "."))
        except ValueError:
            parsed = None
        if parsed is not None:
            return {
                "value_numeric": parsed / 100, "chart_value": parsed,
                "display_value": value.strip(), "value_kind": "percentage_text",
                "validation_status": "valid",
            }
    return {
        "value_numeric": None, "chart_value": None, "display_value": str(value),
        "value_kind": "text", "validation_status": "valid" if metric == "Ghi chú" else "warning",
    }


def _find_header(ws, config: ParserConfig) -> tuple[int, int, int] | None:
    for row_idx in range(1, min(ws.max_row, config.header_search_rows) + 1):
        values = {clean_text(ws.cell(row_idx, col).value).casefold(): col for col in range(1, ws.max_column + 1)}
        project_col = values.get(config.project_column.casefold())
        unit_col = values.get(config.unit_column.casefold())
        if project_col:
            return row_idx, project_col, unit_col or project_col + 1
    return None


def _date_groups(ws, header_row: int, config: ParserConfig, epoch: datetime) -> list[dict[str, Any]]:
    meaningful_headers: list[tuple[int, str]] = []
    for col in range(1, ws.max_column + 1):
        text = clean_text(ws.cell(header_row, col).value)
        if text:
            meaningful_headers.append((col, text))
    groups: list[dict[str, Any]] = []
    for index, (start_col, header) in enumerate(meaningful_headers):
        folded = fold_text(header)
        if not any(fold_text(pattern) in folded for pattern in config.result_header_patterns):
            continue
        parsed_date = extract_date(header, epoch)
        if parsed_date is None:
            continue
        end_col = (meaningful_headers[index + 1][0] - 1) if index + 1 < len(meaningful_headers) else ws.max_column
        groups.append({"date": parsed_date, "start_col": start_col, "end_col": end_col, "header": header})
    return groups


def parse_workbook(source: ExcelSource, config: ParserConfig | None = None) -> ParseResult:
    config = config or ParserConfig()
    records: list[dict[str, Any]] = []
    issues: list[ValidationIssue] = []
    entity_records: list[dict[str, Any]] = []
    detected_dates: set[pd.Timestamp] = set()
    detected_metrics: set[str] = set()

    for ws in source.workbook.worksheets:
        header = _find_header(ws, config)
        if header is None:
            issues.append(ValidationIssue("warning", "HEADER_NOT_FOUND", "Không tìm thấy cột Dự án.", ws.title))
            continue
        header_row, project_col, unit_col = header
        metric_row = header_row + 1
        groups = _date_groups(ws, header_row, config, source.workbook.epoch)
        if not groups:
            issues.append(ValidationIssue("error", "DATE_GROUP_NOT_FOUND", "Không tìm thấy block Kết quả triển khai.", ws.title))
            continue

        tree = EntityTreeBuilder(ws.title, config)
        warned_unknown_units: set[str] = set()
        for row_idx in range(metric_row + 1, ws.max_row + 1):
            label = clean_text(ws.cell(row_idx, project_col).value)
            if not label:
                continue
            sequence = ws.cell(row_idx, 1).value
            unit_value = clean_text(ws.cell(row_idx, unit_col).value)
            node = tree.add_row(row_idx, sequence, label, unit_value)
            if node is None:
                issues.append(ValidationIssue(
                    "error", "ENTITY_WITHOUT_PROJECT", f"Entity không xác định được Project: {label}",
                    ws.title, ws.cell(row_idx, project_col).coordinate,
                ))
                continue
            if node.parser_rule == "fallback_entity":
                issues.append(ValidationIssue(
                    "warning", "FALLBACK_ENTITY_CLASSIFICATION",
                    f"Entity được phân loại bằng fallback rule: {node.entity_label}",
                    ws.title, ws.cell(row_idx, project_col).coordinate,
                ))

            section = _ancestor_label(node, "section", tree.by_id)
            legacy_item = node.entity_label if node.entity_level not in {"project", "section"} else None
            for group in groups:
                detected_dates.add(group["date"])
                for col_idx in range(group["start_col"], group["end_col"] + 1):
                    metric_original = clean_text(ws.cell(metric_row, col_idx).value)
                    if not metric_original:
                        continue
                    cell = ws.cell(row_idx, col_idx)
                    if _is_blank(cell.value, config):
                        continue
                    metric = normalize_metric(metric_original, config)
                    detected_metrics.add(metric)
                    fields = _value_fields(cell, metric, config)
                    record = {
                        "source_file": source.source_file,
                        "source_hash": source.source_hash,
                        "sheet_name": ws.title,
                        "cell_address": cell.coordinate,
                        "source_row": row_idx,
                        "project_id": node.project_id,
                        "project_label": node.project_label,
                        "entity_id": node.entity_id,
                        "parent_entity_id": node.parent_entity_id,
                        "entity_level": node.entity_level,
                        "entity_depth": node.entity_depth,
                        "entity_label": node.entity_label,
                        "entity_path": node.entity_path,
                        "entity_key": node.entity_id,
                        "unit_raw": node.unit_raw,
                        "unit_original": node.unit_original,
                        "unit_normalized": node.unit_normalized,
                        "effective_unit": node.effective_unit,
                        "unit_source_level": node.unit_source_level,
                        "unit_source_entity_id": node.unit_source_entity_id,
                        "parser_rule": node.parser_rule,
                        "parser_confidence": node.parser_confidence,
                        # Compatibility projection for existing exports/consumers.
                        "project": node.project_label,
                        "section": section,
                        "item": legacy_item,
                        "unit": node.effective_unit,
                        "date": group["date"],
                        "metric_original": metric_original,
                        "metric_normalized": metric,
                        "raw_value": cell.value,
                        "number_format": cell.number_format,
                        **fields,
                    }
                    records.append(record)
                    if fields["chart_value"] is not None and node.effective_unit is None and node.entity_id not in warned_unknown_units:
                        warned_unknown_units.add(node.entity_id)
                        issues.append(ValidationIssue(
                            "warning", "UNKNOWN_UNIT",
                            f"Entity có metric dạng số nhưng chưa xác định được unit: {node.entity_path}",
                            ws.title, cell.coordinate,
                        ))
                    if fields["validation_status"] == "warning":
                        issue_code = "MISSING_MARKER" if fields["value_kind"] == "missing_marker" else "NON_NUMERIC_METRIC"
                        issues.append(ValidationIssue(
                            "warning", issue_code,
                            f"Metric '{metric_original}' chứa giá trị không dùng để vẽ: {fields['display_value']}",
                            ws.title, cell.coordinate,
                        ))

        entity_records.extend(node.as_dict() for node in tree.nodes)

    data = pd.DataFrame.from_records(records, columns=OUTPUT_COLUMNS)
    entities = pd.DataFrame.from_records(entity_records, columns=ENTITY_COLUMNS)
    project_nodes = entities[entities["entity_level"] == "project"] if not entities.empty else entities
    section_nodes = entities[entities["entity_level"] == "section"] if not entities.empty else entities
    item_nodes = entities[entities["entity_level"].isin(["item", "subitem"])] if not entities.empty else entities
    manifest = {
        "source_file": source.source_file,
        "source_hash": source.source_hash,
        "sheet_count": len(source.workbook.sheetnames),
        "record_count": len(data),
        "chartable_record_count": int(data["chart_value"].notna().sum()) if not data.empty else 0,
        "project_count": len(project_nodes),
        "section_count": len(section_nodes),
        "item_count": len(item_nodes),
        "entity_count": len(entities),
        "max_entity_depth": int(entities["entity_depth"].max()) if not entities.empty else 0,
        "fallback_entity_count": int((entities["parser_rule"] == "fallback_entity").sum()) if not entities.empty else 0,
        "unknown_unit_count": int(entities["effective_unit"].isna().sum()) if not entities.empty else 0,
        "unit_count": int(entities["unit_normalized"].nunique()) if not entities.empty else 0,
        "date_count": len(detected_dates),
        "metric_count": len(detected_metrics),
        "projects": sorted(project_nodes["project_label"].unique()) if not entities.empty else [],
        "metrics": sorted(detected_metrics),
    }
    return ParseResult(data=data, entities=entities, issues=issues, manifest=manifest)

"""Local, deterministic PDF/DOCX renderers. No remote URLs or browser screenshots."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
import os
from xml.sax.saxutils import escape

from .composer import CALC_LABELS, METRIC_LABELS

RENDERER_VERSION = 'cx-report-renderer-v4'
CONTENT_TYPES = {'pdf': 'application/pdf', 'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'}


def font_paths() -> tuple[Path, Path]:
    configured = os.environ.get('EVP_REPORT_FONT')
    candidates = ([Path(configured)] if configured else []) + [
        Path('C:/Windows/Fonts/segoeui.ttf'), Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')]
    regular = next((p for p in candidates if p.is_file()), None)
    if not regular:
        raise ValueError('Chưa có font Unicode để xuất tiếng Việt. Cấu hình EVP_REPORT_FONT tới font TrueType hợp lệ rồi thử lại.')
    bold = next((p for p in [regular.with_name('segoeuib.ttf'), regular.with_name('DejaVuSans-Bold.ttf')] if p.is_file()), regular)
    return regular, bold


def finding_caption(finding: dict, document: dict | None = None) -> str:
    groups = {}
    for anchor in finding['anchors']:
        groups.setdefault((anchor['entityRef'], anchor['metricCode']), []).append(anchor)
    clauses = []
    multiple_entities = len({a['entityRef'] for a in finding['anchors']}) > 1
    for (entity, code), anchors in groups.items():
        anchors.sort(key=lambda a: a['periodStart'])
        first = anchors[0]
        label = METRIC_LABELS[code]
        if multiple_entities and document:
            chart = next(c for c in document['charts'] if c['chartId'] == first['chartId'])
            label = f"{chart['entityLabel']} · {label}"
        clause = label + ': ' + ' → '.join(f"{a['displayValue']} ({a['periodLabel']})" for a in anchors)
        clauses.append(clause)
    return '; '.join(clauses) + '.'


def markers_for(document: dict, chart: dict) -> list[dict]:
    lookup = {f['findingId']: f for f in document['findings']}
    marks = []
    for number, fid in enumerate(document['selectedFindingIds'], 1):
        for anchor in lookup[fid]['anchors']:
            if anchor['chartId'] == chart['chartId']:
                marks.append({**anchor, 'number': number, 'findingId': fid})
    return marks


def chart_png(document: dict, chart: dict) -> bytes:
    if chart.get('seriesCharts'):
        return panel_png(document, chart)
    from PIL import Image, ImageDraw, ImageFont
    regular, bold = font_paths()
    image = Image.new('RGB', (1600, 680), 'white')
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(str(regular), 23)
    small = ImageFont.truetype(str(regular), 21)
    heading = ImageFont.truetype(str(bold), 27)
    draw.text((70, 24), chart['metricLabel'], font=heading, fill='#233044')
    draw.text((70, 68), f"{chart['calculationLabel']} · {chart['unit']}", font=font, fill='#526176')
    rows = chart['points']
    valid = [p for p in rows if p['value'] is not None]
    if not valid:
        draw.text((70, 200), 'Chưa có dữ liệu.', font=font, fill='#526176')
    else:
        # Same canonical points and selection as preview. Nulls break the line.
        left, right, top, bottom = 145, 1530, 138, 552
        low, high = min(0, min(p['value'] for p in valid)), max(p['value'] for p in valid)
        span = high - low or max(abs(high), 1)
        high += span * .16
        low -= span * .05 if low < 0 else 0
        def y(value):
            return bottom - (value - low) / (high - low) * (bottom - top)
        # Use actual dates so missing intervals and unequal period lengths stay visible.
        from datetime import date
        dates = [date.fromisoformat(p['periodStart']).toordinal() for p in rows]
        first_date, last_date = min(dates), max(dates)
        def x(index):
            return (left + right) / 2 if first_date == last_date else left + (dates[index] - first_date) / (last_date - first_date) * (right - left)
        for step in range(5):
            value = low + (high - low) * step / 4
            ypos = y(value)
            draw.line((left, ypos, right, ypos), fill='#e8ecf1', width=1)
            label = f'{value:,.2f}'.rstrip('0').rstrip('.').replace(',', ' ')
            draw.text((left - 16, ypos), label, font=small, fill='#526176', anchor='rm')
        draw.line((left, top, left, bottom), fill='#8a95a6', width=2)
        draw.line((left, bottom, right, bottom), fill='#8a95a6', width=2)
        point_lookup = {p['factId']: index for index, p in enumerate(rows) if p.get('factId')}
        marks = markers_for(document, chart)
        intervals = {}
        for mark in marks:
            index = point_lookup.get(mark['factId'])
            if index is not None:
                intervals.setdefault(mark['number'], []).append(index)
        for indices in intervals.values():
            if len(indices) > 1 and not any(rows[i]['value'] is None for i in range(min(indices), max(indices)+1)):
                draw.rectangle((x(min(indices)), top, x(max(indices)), bottom), fill='#f5f3fb')
        previous = None
        for index, row in enumerate(rows):
            if row['value'] is None:
                previous = None
                continue
            point = (x(index), y(row['value']))
            if previous:
                draw.line((*previous, *point), fill='#315b83', width=4)
            draw.ellipse((point[0]-5, point[1]-5, point[0]+5, point[1]+5), fill='#315b83')
            previous = point
        by_point = {}
        for mark in marks:
            index = point_lookup.get(mark['factId'])
            if index is not None:
                by_point.setdefault(index, []).append(mark['number'])
        for index, numbers in by_point.items():
            px, py = x(index), y(rows[index]['value'])
            label = ','.join(map(str, dict.fromkeys(numbers)))
            draw.rounded_rectangle((px-26, py-52, px+26, py-15), radius=6, fill='#6253b5')
            draw.text((px, py-34), label, font=small, fill='white', anchor='mm')
        indices = sorted(set(round(i * (len(rows)-1) / max(1, min(len(rows), 5)-1)) for i in range(min(len(rows), 5))))
        for index in indices:
            row = rows[index]
            # Preserve the canonical grouping label (week/month/quarter), not
            # a date that makes a weekly report look like daily observations.
            label = row['periodLabel']
            half_width = draw.textlength(label, font=small) / 2
            label_x = min(1580-half_width, max(20+half_width, x(index)))
            draw.text((label_x, bottom+22), label, font=small, fill='#526176', anchor='mt')
        draw.text((left, 625), 'Khoảng tô nhạt: giai đoạn được đánh dấu. Số trên biểu đồ khớp mục Điểm đáng chú ý.', font=small, fill='#526176')
    stream = BytesIO()
    from PIL.PngImagePlugin import PngInfo
    provenance = PngInfo()
    provenance.add_text('Origin', 'Deterministic Analytics Engine series; locally drawn by Pillow, not an AI image or dashboard screenshot.')
    provenance.add_text('Report', f"{document['reportId']} / revision {document['revision']} / {chart['chartId']}")
    provenance.add_text('SourceChecksum', document['dataAsOf']['checksum'])
    image.save(stream, 'PNG', pnginfo=provenance)
    return stream.getvalue()


def panel_png(document: dict, panel: dict) -> bytes:
    """Same typed series/color/axis contract as preview, drawn locally for export."""
    from PIL import Image, ImageDraw, ImageFont
    from PIL.PngImagePlugin import PngInfo
    from .story import chart_style, panel_figure
    charts = panel['seriesCharts']
    figure = panel_figure(charts)
    regular, bold = font_paths()
    font = ImageFont.truetype(str(regular), 20)
    small = ImageFont.truetype(str(regular), 18)
    image = Image.new('RGB', (1600, 760), 'white')
    draw = ImageDraw.Draw(image)
    units = list(dict.fromkeys(c['unit'] for c in charts))
    dates = figure['layout']['xaxis']['categoryarray']
    left, right, top, bottom = 145, 1450, 190, 615
    draw.text((left, top-28), units[0], font=small, fill='#526176')
    if len(units) == 2:
        draw.text((right, top-28), units[1], font=small, fill='#526176', anchor='rt')
    scales = {}
    for unit in units:
        values = [p['value'] for c in charts if c['unit'] == unit for p in c['points'] if p['value'] is not None]
        low, high = min([0, *values]), max([0, *values])
        span = high-low or 1
        scales[unit] = (low, high+span*.2)
    def x(day):
        return (left+right)/2 if len(dates) < 2 else left+(dates.index(day)+.5)/len(dates)*(right-left)
    def y(value, unit):
        low, high = scales[unit]
        return bottom-(value-low)/(high-low)*(bottom-top)
    for unit_index, unit in enumerate(units):
        low, high = scales[unit]
        for step in range(5):
            value = low+(high-low)*step/4
            ypos = y(value,unit)
            if unit_index == 0:
                draw.line((left,ypos,right,ypos), fill='#eef0f4')
            draw.text((left-12 if unit_index == 0 else right+12,ypos), f'{value:,.2f}'.rstrip('0').rstrip('.'),
                      font=small,fill='#526176',anchor='rm' if unit_index == 0 else 'lm')
    for index, chart in enumerate(charts):
        kind, color = chart_style(chart)
        label = figure['data'][index]['name']
        lx, ly = 65+(index%2)*760, 68+(index//2)*36
        draw.rectangle((lx,ly+6,lx+18,ly+20),fill=color)
        draw.text((lx+30,ly),label,font=font,fill='#526176')
        previous = None
        bar_width = min(100,(right-left)/max(len(dates),1)*.7)
        for point in chart['points']:
            if point['value'] is None:
                previous = None
                continue
            px, py = x(point['periodStart']), y(point['value'],chart['unit'])
            if kind == 'bar':
                draw.rectangle((px-bar_width/2,min(py,y(0,chart['unit'])),px+bar_width/2,max(py,y(0,chart['unit']))),fill=color)
            else:
                if previous:
                    draw.line((*previous,px,py),fill=color,width=4)
                draw.ellipse((px-5,py-5,px+5,py+5),fill=color)
            previous = (px,py)
    by_period = {}
    for chart in charts:
        for mark in markers_for(document,chart):
            point = next((p for p in chart['points'] if p['factId'] == mark['factId']), None)
            if point and point['value'] is not None:
                by_period.setdefault(point['periodStart'],set()).add(mark['number'])
    for date, numbers in by_period.items():
        px, py = x(date), top-8
        label = ','.join(map(str,sorted(numbers)))
        half = max(18,draw.textlength(label,font=small)/2+6)
        draw.rectangle((px-half,py-27,px+half,py),fill='#6253b5')
        draw.text((px,py-14),label,font=small,fill='white',anchor='mm')
    count = min(len(dates),5)
    indices = sorted(set(round(i*(len(dates)-1)/max(1,count-1)) for i in range(count)))
    for index in indices:
        label = figure['layout']['xaxis']['ticktext'][index]
        half = draw.textlength(label,font=small)/2
        draw.text((min(1580-half,max(20+half,x(dates[index]))),bottom+20),label,font=small,fill='#526176',anchor='mt')
    draw.text((65,717),'Số đánh dấu khớp mục Điểm đáng chú ý. Mỗi trục giữ đơn vị được ghi trên biểu đồ.',font=small,fill='#526176')
    stream = BytesIO()
    metadata = PngInfo()
    metadata.add_text('Origin','Deterministic report snapshot; shared story figure contract; locally drawn with Pillow.')
    metadata.add_text('Report',f"{document['reportId']} / revision {document['revision']}")
    metadata.add_text('SourceChecksum',document['dataAsOf']['checksum'])
    image.save(stream,'PNG',pnginfo=metadata)
    return stream.getvalue()


def prose_rows(blocks: list[dict]) -> list[tuple[str, object]]:
    """Keep AI, Engine and manual authorship visible without repeating a tag per sentence."""
    labels = {'ai':'AI đã kiểm chứng', 'manual':'Người dùng chỉnh sửa — đã kiểm chứng',
              'mixed':'Tổng hợp từ AI và số liệu đã kiểm chứng', 'deterministic':'Tổng hợp từ số liệu'}
    rows, previous = [], None
    for block in blocks:
        source = block['source']
        if source != previous:
            rows.append(('note', 'Nguồn diễn giải: ' + labels[source]))
        rows.append(('paragraph', block['text']))
        previous = source
    return rows


def sections(document: dict) -> list[tuple[str, list[tuple[str, object]]]]:
    ctx, window, data = document['context'], document['window'], document['dataAsOf']
    members = list(dict.fromkeys(c['entityLabel'] for c in document['charts']))
    metadata = [f"{document['title']} — BẢN NHÁP / DRAFT",
        f"Dự án: {ctx['project']} · Phiên bản: {document['revision']}",
        f"Thời gian dữ liệu: {window['start']} → {window['end']} · Nhóm kỳ: { {'day':'Ngày','week':'Tuần','month':'Tháng','quarter':'Quý'}[window['groupBy']]}",
        f"Nguồn phân tích: {'Thống kê' if ctx['view'] == 'statistics' else 'Tổng quan'} · Cách tính: {CALC_LABELS.get(ctx['calculation'], 'Tổng và trung bình/ngày')}",
        'Phạm vi: ' + '; '.join(members),
        f"Dữ liệu nguồn: {data['committedImportRef']} · Chụp dữ liệu: {data['generatedAt']}",
        f"Tạo báo cáo: {document['createdAt']} · Template: {document['template']['id']} {document['template']['version']}",
        'Kiểm tra bản nháp tại lần xuất: ' + ('Đã kiểm tra · ' + str(document.get('reviewAtExport', {}).get('checkedAt'))
            if document.get('reviewAtExport', {}).get('status') == 'checked' else 'Chưa đánh dấu đã kiểm tra'),
        'Đây là bản nháp. Xác nhận đã kiểm tra không phải phê duyệt có danh tính xác thực.']
    summary = prose_rows(document['executiveSummary'])
    kpis = [['Vấn đề / KPI / cách tính', 'Mức ở kỳ có dữ liệu cuối', 'Đơn vị', 'Số kỳ có dữ liệu']]
    for kpi in document['kpis']:
        kpis.append([f"{kpi['entityLabel']} / {kpi['metricLabel']} / {kpi['calculationLabel']}",
                     f"{kpi['displayValue']} ({kpi['periodLabel']})", kpi['unit'],
                     f"{kpi['quality']['validPeriodCount']}/{kpi['quality']['expectedPeriodCount']}"])
    story = []
    keys = list(dict.fromkeys((c['entityRef'], c['calculation']) for c in document['charts']))
    if document['template']['version'] != '1.0':
        from .story import story_panels
        for panel in document.get('storyPanels') or story_panels(document):
            charts = [c for c in document['charts'] if c['chartId'] in panel['chartIds']]
            story.append(('subheading', panel['entityLabel'] + (' · ' + panel['title'] if panel.get('title') else '')))
            story.append(('chart',{'metricLabel':panel['entityLabel'],'calculationLabel': ' / '.join(dict.fromkeys(c['calculationLabel'] for c in charts)),
                                   'seriesCharts':charts}))
            previous_calculation = None
            previous_source = None
            lookup = {b['blockId']: b for b in document['blocks']}
            for block in [*[lookup[bid] for bid in panel['blockIds'] if bid in lookup], *panel.get('readings', [])]:
                calculation = block.get('calculation') or 'both'
                if ctx['calculation'] == 'both' and calculation != previous_calculation:
                    story.append(('note',CALC_LABELS.get(calculation, 'Tổng và trung bình/ngày')))
                previous_calculation = calculation
                rows = prose_rows([block])
                if document['template']['version'] == '1.3' and block['source'] == previous_source:
                    rows = [row for row in rows if row[0] != 'note']
                story += rows
                previous_source = block['source']
        keys = []
    for entity, calculation in keys:
        charts = [c for c in document['charts'] if c['entityRef'] == entity and c['calculation'] == calculation]
        story.append(('subheading', f"{charts[0]['entityLabel']} · {CALC_LABELS[calculation]}"))
        paragraphs = [b for b in document['blocks'] if b['section'] != 'overview' and b.get('entityRef') == entity and b.get('calculation') == calculation]
        for chart in charts:
            story.append(('chart', chart))
        story += prose_rows(paragraphs)
    shared = [b for b in document['blocks'] if b['section'] == 'group_relationships']
    if shared:
        story.append(('subheading','Liên hệ giữa các vấn đề'))
        story += prose_rows(shared)
    highlights = []
    lookup = {f['findingId']: f for f in document['findings']}
    for number, fid in enumerate(document['selectedFindingIds'], 1):
        finding = lookup[fid]
        highlights.append(('subheading', f"{number}. {finding['entityLabel']} · {finding['title']} · {CALC_LABELS[finding['calculation']]}"))
        # Do not copy the whole phase template again. Pointing is an indexed data caption.
        highlights.append(('paragraph', finding_caption(finding, document)))
    if not highlights:
        highlights = [('paragraph', 'Chưa chọn điểm đánh dấu. Phần diễn biến và dữ liệu vẫn được giữ đầy đủ.')]
    if document['limitations']:
        highlights.append(('subheading', 'Giới hạn khi đọc báo cáo'))
        highlights += [('note', note) for note in document['limitations']]
    if ctx['excluded']:
        highlights += [('note', f"Không có dữ liệu để phân tích: {e['entityLabel']} ({e['reason']}).") for e in ctx['excluded']]
    if document['userNotes']:
        highlights.extend([('subheading', 'Ghi chú của người dùng — chưa kiểm chứng'), ('paragraph', document['userNotes'])])
    highlights.append(('subheading', 'Nguồn và phiên bản phân tích'))
    highlights.extend([('note', f"Report ID: {document['reportId']} · Revision {document['revision']} · Snapshot: {data['snapshotId']}"),
        ('note', f"Checksum dữ liệu: {data['checksum']}"),
        ('note', f"Prompt: {ctx['promptVersion']} · Policy: {ctx['policyVersion']} · Model: {document['generation']['provider']['model']}"),
        ('note', 'DOCX chỉnh sửa ngoài ứng dụng không còn được đảm bảo kiểm chứng như phiên bản đã xuất.')])
    # A source appendix with exact period/entity/metric identity, not arbitrary URLs.
    for chart in document['charts']:
        for point in chart['points']:
            if point.get('evidenceId'):
                target = next(e['target'] for e in document['evidence'] if e['evidenceId'] == point['evidenceId'])
                ref = target.get('aggregateRef') or f"{target['observationRef']} / {target['lineageRef']}"
                highlights.append(('note', f"{chart['entityLabel']} · {chart['metricLabel']} · {chart['calculationLabel']} · {point['periodLabel']}: {point['displayValue']} · {ref}"))
    return [('Thông tin báo cáo', [('paragraph', text) for text in metadata]), ('Tóm tắt điều hành', summary),
            ('Tổng quan KPI', [('note', 'Bảng thể hiện mức tại kỳ có dữ liệu cuối, không phải tổng cộng của toàn bộ chuỗi.'), ('table', kpis)]),
            ('Diễn biến trong kỳ', story), ('Điểm đáng chú ý', highlights)]


def render_pdf(document: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
    regular, bold = font_paths()
    pdfmetrics.registerFont(TTFont('CXReport', str(regular)))
    pdfmetrics.registerFont(TTFont('CXReportBold', str(bold)))
    body = ParagraphStyle('body', fontName='CXReport', fontSize=10, leading=15, spaceAfter=9, textColor=colors.HexColor('#38465a'), splitLongWords=True)
    heading = ParagraphStyle('heading', parent=body, fontName='CXReportBold', fontSize=15, leading=20, spaceBefore=18, spaceAfter=10, keepWithNext=True)
    sub = ParagraphStyle('sub', parent=body, fontName='CXReportBold', fontSize=11, leading=16, spaceBefore=10, keepWithNext=True)
    note = ParagraphStyle('note', parent=body, fontSize=8, leading=12)
    stream = BytesIO()
    pdf = SimpleDocTemplate(stream, pagesize=A4, leftMargin=38, rightMargin=38, topMargin=42, bottomMargin=42,
                           title=document['title'], author='Automated CX Report')
    flow = []
    def p(text, style=body):
        return Paragraph(escape(str(text)).replace('\n', '<br/>'), style)
    for title, content in sections(document):
        flow.append(p(title, heading))
        for kind, value in content:
            if kind == 'chart':
                if not value.get('seriesCharts'):
                    flow.append(p(f"{value['metricLabel']} · {value['calculationLabel']}", sub))
                flow += [Image(BytesIO(chart_png(document, value)), width=519, height=246 if value.get('seriesCharts') else 221), Spacer(1, 10)]
            elif kind == 'table':
                table = Table([[p(cell, note) for cell in row] for row in value], colWidths=[210, 155, 75, 79], repeatRows=1, hAlign='LEFT')
                table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#f0eef8')), ('VALIGN',(0,0),(-1,-1),'TOP'),
                    ('LINEBELOW',(0,0),(-1,0),1,colors.HexColor('#dde2e8')), ('BOTTOMPADDING',(0,0),(-1,-1),8), ('TOPPADDING',(0,0),(-1,-1),8)]))
                flow.append(table)
            else:
                flow.append(p(value, sub if kind == 'subheading' else note if kind == 'note' else body))
    def footer(canvas, doc):
        canvas.setFont('CXReport', 8)
        canvas.setFillColor(colors.HexColor('#526176'))
        canvas.drawString(38, 24, f"Automated CX Report · DRAFT · Phiên bản {document['revision']}")
        canvas.drawRightString(A4[0]-38, 24, f'Trang {doc.page}')
    pdf.build(flow, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()


def render_docx(document: dict) -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.65)
    normal = doc.styles['Normal']
    normal.font.name = 'Segoe UI'
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string('38465A')
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.25
    doc.core_properties.title = document['title']
    doc.core_properties.author = 'Automated CX Report'
    doc.core_properties.subject = f"DRAFT · {document['reportId']} · revision {document['revision']}"
    for title, content in sections(document):
        doc.add_heading(title, level=1)
        for kind, value in content:
            if kind == 'chart':
                if not value.get('seriesCharts'):
                    doc.add_heading(f"{value['metricLabel']} · {value['calculationLabel']}", level=3)
                doc.add_picture(BytesIO(chart_png(document, value)), width=Inches(6.5))
            elif kind == 'table':
                table = doc.add_table(rows=1, cols=len(value[0]))
                table.style = 'Light Shading Accent 1'
                for index, cell in enumerate(value[0]):
                    table.rows[0].cells[index].text = cell
                # Repeat the heading on long tables in Word, not just the first page.
                from docx.oxml import OxmlElement
                tr_props = table.rows[0]._tr.get_or_add_trPr()
                tr_props.append(OxmlElement('w:tblHeader'))
                for row in value[1:]:
                    cells = table.add_row().cells
                    for index, cell in enumerate(row):
                        cells[index].text = str(cell)
            elif kind == 'subheading':
                doc.add_heading(str(value), level=2)
            else:
                doc.add_paragraph(str(value))
    section.footer.paragraphs[0].text = f"Automated CX Report · BẢN NHÁP / DRAFT · Phiên bản {document['revision']}"
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def render(document: dict, format: str) -> bytes:
    if format not in CONTENT_TYPES:
        raise ValueError('Chỉ hỗ trợ PDF và DOCX trong phiên bản này.')
    return render_pdf(document) if format == 'pdf' else render_docx(document)

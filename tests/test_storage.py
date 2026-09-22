from __future__ import annotations

import sqlite3
from contextlib import closing

from openpyxl import load_workbook

from excel_visualization_pipeline.storage import (
    ImportScope,
    backup_database,
    import_workbook,
    initialize_database,
    load_current_data,
    load_current_entities,
    load_import_history,
    verify_database,
)


def test_initializes_schema_idempotently(storage_workspace):
    database = storage_workspace / "analytics.sqlite3"

    initialize_database(database)
    initialize_database(database)

    report = verify_database(database)
    assert report["is_valid"] is True
    assert report["migration_count"] == 6
    assert report["latest_committed_run_id"] is None


def test_migration_six_repairs_early_aggregate_schema(storage_workspace):
    database = storage_workspace / "analytics.sqlite3"
    initialize_database(database)
    with closing(sqlite3.connect(database)) as connection:
        connection.execute("DROP TABLE aggregate_cursor_secret")
        connection.execute("DELETE FROM schema_migrations WHERE version = 6")
        connection.commit()

    initialize_database(database)
    with closing(sqlite3.connect(database)) as connection:
        assert connection.execute("SELECT length(secret) FROM aggregate_cursor_secret").fetchone()[0] == 32
        assert connection.execute("SELECT COUNT(*) FROM schema_migrations WHERE version = 6").fetchone()[0] == 1


def test_commits_and_skips_duplicate_without_duplicating_data(storage_workspace, sample_workbook):
    database = storage_workspace / "analytics.sqlite3"

    first = import_workbook(database, sample_workbook, source_key="sample")
    second = import_workbook(database, sample_workbook, source_key="sample")

    assert first.outcome.status == "committed"
    assert first.outcome.inserted_count == 6
    assert second.outcome.status == "duplicate"
    assert second.outcome.duplicate_of_run_id == first.outcome.run_id
    assert len(load_current_data(database, "sample")) == 6
    assert len(load_current_entities(database, "sample")) == 3
    history = load_import_history(database, "sample")
    assert history["attempt_status"].tolist() == ["duplicate", "committed"]
    with closing(sqlite3.connect(database)) as connection:
        assert connection.execute("SELECT COUNT(*) FROM import_runs").fetchone()[0] == 1
        assert connection.execute("SELECT COUNT(*) FROM observation_revisions").fetchone()[0] == 6
        assert connection.execute("SELECT COUNT(*) FROM import_observation_presence").fetchone()[0] == 6
        assert connection.execute("SELECT COUNT(*) FROM observation_public_refs").fetchone()[0] == 6
        assert connection.execute("SELECT COUNT(*) FROM observation_lineage_snapshots").fetchone()[0] == 6


def test_incremental_updates_history_and_preserves_missing_dates(storage_workspace, sample_workbook):
    database = storage_workspace / "analytics.sqlite3"
    first = import_workbook(database, sample_workbook, source_key="sample")
    assert first.outcome.status == "committed"
    initial = load_current_data(database, "sample")
    initial_refs = {
        (row.entity_id, row.date.date().isoformat(), row.metric_code): row.observation_ref
        for row in initial.itertuples(index=False)
    }

    workbook = load_workbook(sample_workbook)
    sheet = workbook.active
    sheet["H7"] = 7
    sheet.unmerge_cells("D3:F3")
    for row in (3, 4):
        for column in range(4, 7):
            sheet.cell(row, column).value = None
    workbook.save(sample_workbook)

    second = import_workbook(
        database,
        sample_workbook,
        source_key="sample",
        mode="incremental",
    )

    assert second.outcome.status == "committed"
    assert second.outcome.updated_count == 1
    assert second.outcome.unchanged_count == 2
    assert second.outcome.deleted_count == 0
    current = load_current_data(database, "sample")
    assert len(current) == 6
    current_refs = {
        (row.entity_id, row.date.date().isoformat(), row.metric_code): row.observation_ref
        for row in current.itertuples(index=False)
    }
    assert current_refs == initial_refs
    error = current[
        (current["date"].dt.strftime("%Y-%m-%d") == "2026-09-13")
        & (current["metric_normalized"] == "Báo sai/Lỗi")
    ].iloc[0]
    assert error["chart_value"] == 7
    with closing(sqlite3.connect(database)) as connection:
        revisions = connection.execute(
            """
            SELECT COUNT(*)
            FROM observation_revisions r
            JOIN observations o ON o.observation_id = r.observation_id
            WHERE o.observed_date = '2026-09-13' AND o.metric_code = 'error'
            """
        ).fetchone()[0]
    assert revisions == 2
    with closing(sqlite3.connect(database)) as connection:
        observation_refs = connection.execute(
            "SELECT COUNT(DISTINCT observation_ref) FROM observation_public_refs"
        ).fetchone()[0]
        lineage_refs = connection.execute(
            "SELECT COUNT(DISTINCT lineage_ref) FROM observation_lineage_snapshots"
        ).fetchone()[0]
    assert observation_refs == 6
    assert lineage_refs == 9


def test_lineage_only_change_does_not_create_business_revision(storage_workspace, sample_workbook):
    database = storage_workspace / "analytics.sqlite3"
    import_workbook(database, sample_workbook, source_key="sample")

    workbook = load_workbook(sample_workbook)
    workbook.active.insert_rows(7)
    workbook.save(sample_workbook)

    second = import_workbook(database, sample_workbook, source_key="sample")

    assert second.outcome.status == "committed"
    assert second.outcome.updated_count == 0
    assert second.outcome.unchanged_count == 6
    assert second.outcome.lineage_changed_count == 6
    with closing(sqlite3.connect(database)) as connection:
        assert connection.execute("SELECT COUNT(*) FROM observation_revisions").fetchone()[0] == 6
        assert connection.execute("SELECT COUNT(*) FROM import_observation_presence").fetchone()[0] == 12


def test_full_snapshot_tombstones_only_declared_scope(storage_workspace, sample_workbook):
    database = storage_workspace / "analytics.sqlite3"
    import_workbook(database, sample_workbook, source_key="sample")

    workbook = load_workbook(sample_workbook)
    sheet = workbook.active
    sheet.unmerge_cells("D3:F3")
    for row in (3, 4):
        for column in range(4, 7):
            sheet.cell(row, column).value = None
    workbook.save(sample_workbook)

    second = import_workbook(
        database,
        sample_workbook,
        source_key="sample",
        mode="full_snapshot",
        scopes=[
            ImportScope(
                sheet_name="Data",
                date_from="2026-09-12",
                date_to="2026-09-12",
                missing_policy="tombstone",
            )
        ],
    )

    assert second.outcome.status == "committed"
    assert second.outcome.deleted_count == 3
    current = load_current_data(database, "sample")
    assert len(current) == 3
    assert set(current["date"].dt.strftime("%Y-%m-%d")) == {"2026-09-13"}


def test_rejected_import_is_audited_without_changing_current_data(storage_workspace, sample_workbook):
    database = storage_workspace / "analytics.sqlite3"
    import_workbook(database, sample_workbook, source_key="sample")

    workbook = load_workbook(sample_workbook)
    workbook.active["D3"] = "Không còn block ngày"
    workbook.active["G3"] = "Không còn block ngày"
    workbook.save(sample_workbook)

    rejected = import_workbook(database, sample_workbook, source_key="sample")

    assert rejected.outcome.status == "rejected"
    assert len(load_current_data(database, "sample")) == 6
    history = load_import_history(database, "sample")
    assert history.iloc[0]["attempt_status"] == "rejected"
    with closing(sqlite3.connect(database)) as connection:
        assert connection.execute("SELECT COUNT(*) FROM import_runs").fetchone()[0] == 1
        assert connection.execute(
            "SELECT COUNT(*) FROM validation_issues WHERE severity = 'error'"
        ).fetchone()[0] >= 1


def test_same_artifact_cannot_be_silently_reapplied_with_another_mode(
    storage_workspace, sample_workbook
):
    database = storage_workspace / "analytics.sqlite3"

    full = import_workbook(database, sample_workbook, source_key="sample")
    incremental = import_workbook(
        database,
        sample_workbook,
        source_key="sample",
        mode="incremental",
    )

    assert full.outcome.status == "committed"
    assert incremental.outcome.status == "rejected"
    assert "đã được áp dụng" in incremental.outcome.message
    with closing(sqlite3.connect(database)) as connection:
        assert connection.execute("SELECT COUNT(*) FROM import_runs").fetchone()[0] == 1
        issue = connection.execute(
            "SELECT code FROM validation_issues WHERE attempt_id = ? AND severity = 'error'",
            (incremental.outcome.attempt_id,),
        ).fetchone()
        assert issue[0] == "ARTIFACT_ALREADY_APPLIED"


def test_explicit_replay_can_restore_values_from_an_older_artifact(
    storage_workspace, sample_workbook
):
    database = storage_workspace / "analytics.sqlite3"
    original_payload = sample_workbook.read_bytes()
    import_workbook(database, sample_workbook, source_key="sample")

    workbook = load_workbook(sample_workbook)
    workbook.active["H7"] = 7
    workbook.save(sample_workbook)
    changed = import_workbook(database, sample_workbook, source_key="sample")
    assert changed.outcome.updated_count == 1

    restored = import_workbook(
        database,
        original_payload,
        source_key="sample",
        allow_replay=True,
    )

    assert restored.outcome.status == "committed"
    assert restored.outcome.updated_count == 1
    current = load_current_data(database, "sample")
    error = current[
        (current["date"].dt.strftime("%Y-%m-%d") == "2026-09-13")
        & (current["metric_normalized"] == "Báo sai/Lỗi")
    ].iloc[0]
    assert error["chart_value"] == 6


def test_online_backup_is_restorable(storage_workspace, sample_workbook):
    database = storage_workspace / "analytics.sqlite3"
    backup = storage_workspace / "backups" / "analytics.sqlite3"
    import_workbook(database, sample_workbook, source_key="sample")

    backup_database(database, backup)

    assert backup.exists()
    report = verify_database(backup)
    assert report["is_valid"] is True
    assert len(load_current_data(backup, "sample")) == 6

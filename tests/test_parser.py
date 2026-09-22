import pandas as pd

from excel_visualization_pipeline.config import ParserConfig
from excel_visualization_pipeline.ingestion import load_excel
from excel_visualization_pipeline.parser import parse_workbook
from excel_visualization_pipeline.pipeline import run_pipeline


def test_extracts_traceable_long_format(sample_workbook):
    result = run_pipeline(sample_workbook)
    assert result.report.is_valid
    assert result.manifest["project_count"] == 1
    assert result.manifest["date_count"] == 2
    assert result.manifest["observed_date_min"] == "2026-09-12"
    assert result.manifest["observed_date_max"] == "2026-09-13"
    assert len(result.data) == 6
    rate = result.data[result.data["cell_address"] == "F7"].iloc[0]
    assert rate["raw_value"] == 0.08
    assert rate["value_numeric"] == 0.08
    assert rate["chart_value"] == 8.0
    assert rate["display_value"] == "8.00%"
    assert rate["project"] == "Alpha"
    assert rate["section"] == "1.1. Chất lượng cảnh báo"
    assert rate["item"] == "Camera"
    assert rate["entity_depth"] == 2
    assert rate["entity_path"] == "Alpha > 1.1. Chất lượng cảnh báo > Camera"
    assert rate["parent_entity_id"]
    assert rate["effective_unit"] == "Cảnh báo"
    assert rate["unit_source_level"] == "project"


def test_keeps_zero_distinct_from_not_recorded_blank(sample_workbook):
    from openpyxl import load_workbook

    workbook = load_workbook(sample_workbook)
    workbook.active["D7"] = 0
    workbook.active["E7"] = None
    workbook.save(sample_workbook)
    result = run_pipeline(sample_workbook)
    assert "D7" in set(result.data["cell_address"])
    assert "E7" in set(result.data["cell_address"])
    assert result.data.set_index("cell_address").loc["D7", "chart_value"] == 0
    blank = result.data.set_index("cell_address").loc["E7"]
    assert blank["value_kind"] == "not_recorded"
    assert pd.isna(blank["chart_value"])
    assert blank["display_value"] == ""
    assert "% báo sai > 0" in blank["data_note"]
    assert result.manifest["inconsistent_error_metric_count"] == 1
    assert any(
        issue.code == "INCONSISTENT_ERROR_METRICS"
        for issue in result.report.warnings
    )


def test_defaults_blank_rate_only_when_no_error_was_recorded(sample_workbook):
    from openpyxl import load_workbook

    workbook = load_workbook(sample_workbook)
    workbook.active["E7"] = None
    workbook.active["F7"] = None
    workbook.active["I7"] = None
    workbook.save(sample_workbook)

    result = run_pipeline(sample_workbook)
    records = result.data.set_index("cell_address")

    assert records.loc["F7", "value_kind"] == "default_zero_rate"
    assert records.loc["F7", "chart_value"] == 0
    assert records.loc["F7", "display_value"] == "0%"
    assert pd.isna(records.loc["F7", "raw_value"])
    assert records.loc["I7", "value_kind"] == "not_recorded"
    assert pd.isna(records.loc["I7", "chart_value"])


def test_filters_dates_before_configured_minimum(sample_workbook):
    result = parse_workbook(
        load_excel(sample_workbook),
        ParserConfig(minimum_data_date="2026-09-13"),
    )

    assert result.manifest["minimum_data_date"] == "2026-09-13"
    assert result.manifest["observed_date_min"] == "2026-09-13"
    assert result.manifest["observed_date_max"] == "2026-09-13"
    assert result.manifest["date_count"] == 1
    assert len(result.data) == 3
    assert pd.to_datetime(result.data["date"]).min() == pd.Timestamp("2026-09-13")

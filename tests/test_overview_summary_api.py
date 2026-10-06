from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import api
from excel_visualization_pipeline.storage import import_workbook, latest_committed_version
from excel_visualization_pipeline.storage import repository


def test_additive_contract_scope_independence_and_source_validation(monkeypatch, storage_workspace, sample_workbook):
    book = load_workbook(sample_workbook)
    book.active["D5"] = 300
    book.active["G5"] = 500
    book.save(sample_workbook)
    db = storage_workspace / "summary.sqlite3"
    import_workbook(db, sample_workbook, source_key="sample")
    monkeypatch.setattr(api, "DB_PATH", db)
    monkeypatch.setattr(api, "SOURCE_KEY", "sample")
    client = TestClient(api.app)
    first = client.get("/api/projects/Alpha/workspace", params={"view": "overview"})
    assert first.status_code == 200
    payload = first.json()
    summary = payload["overviewSummary"]
    assert summary["metricKey"] == "Tổng số"
    assert summary["scope"] == "project"
    assert summary["issueCount"]["value"] == 1
    assert summary["peak"]["value"] == 500
    assert summary["lowest"]["value"] == 300
    assert summary["largestChange"]["absolute"] == 200
    assert summary["window"] == payload["window"]
    assert payload["dataVersion"]["committedImportRef"].startswith("imp_")
    children = client.get("/api/projects/Alpha/workspace", params={"view": "overview", "scope": "children"})
    assert children.json()["overviewSummary"] == summary
    assert client.get("/api/projects/Alpha/workspace", params={"view": "overview", "overview_source": "other-project"}).status_code == 422
    statistics = client.get("/api/projects/Alpha/workspace", params={"view": "statistics"})
    assert statistics.status_code == 200
    assert statistics.json()["overviewSummary"] is None
    assert statistics.json()["statisticsSummary"]["metricKey"] == "Tổng số"
    assert statistics.json()["statisticsSummary"]["requestedMode"] == "both"
    daily = client.get("/api/projects/Alpha/workspace", params={"view": "statistics", "statistics_group": "day", "statistics_mode": "average"})
    stats_summary = daily.json()["statisticsSummary"]
    assert stats_summary["peak"]["value"] == 500
    assert stats_summary["lowest"]["value"] == 300
    assert stats_summary["calculation"] == "average_per_day"
    assert stats_summary["peak"]["unit"].endswith("/ngày")
    assert client.get("/api/projects/Alpha/workspace", params={"view": "statistics", "overview_source": "other-project"}).status_code == 422
    empty = client.get("/api/projects/Alpha/workspace", params={"view": "statistics", "statistics_group": "quarter", "include_incomplete": False})
    assert empty.status_code == 200
    assert empty.json()["statisticsSummary"]["window"] is None
    assert empty.json()["statisticsSummary"]["peak"]["status"] == "no_data"


def test_read_snapshot_remains_coherent_when_import_commits_between_view_reads(monkeypatch, storage_workspace, sample_workbook):
    db = storage_workspace / "snapshot.sqlite3"
    import_workbook(db, sample_workbook, source_key="sample")
    old_version = latest_committed_version(db, "sample")
    original_read = repository.pd.read_sql_query
    committed = False

    def commit_during_read(query, connection, **kwargs):
        nonlocal committed
        if "v_current_entities" in query and not committed:
            committed = True
            assert connection.in_transaction
            book = load_workbook(sample_workbook)
            book.active["D7"] = 200
            book.save(sample_workbook)
            assert import_workbook(db, sample_workbook, source_key="sample").outcome.status == "committed"
        return original_read(query, connection, **kwargs)

    monkeypatch.setattr(repository.pd, "read_sql_query", commit_during_read)
    version, data, _ = repository.load_current_snapshot(db, "sample")
    assert committed
    assert version == old_version
    assert data.loc[data.metric_code.eq("total")].chart_value.tolist() == [100, 120]
    new_version, new_data, _ = repository.load_current_snapshot(db, "sample")
    assert new_version.run_id > version.run_id
    assert new_data.loc[new_data.metric_code.eq("total")].chart_value.tolist() == [200, 120]


def test_workspace_rekeys_cache_if_committed_version_changes_before_read(monkeypatch, storage_workspace, sample_workbook):
    db = storage_workspace / "retry.sqlite3"
    import_workbook(db, sample_workbook, source_key="sample")
    old_version = latest_committed_version(db, "sample")
    book = load_workbook(sample_workbook)
    book.active["D7"] = 200
    book.save(sample_workbook)
    import_workbook(db, sample_workbook, source_key="sample")
    new_version = latest_committed_version(db, "sample")
    monkeypatch.setattr(api, "DB_PATH", db)
    monkeypatch.setattr(api, "SOURCE_KEY", "sample")
    calls = 0

    def version_at_read(*args):
        nonlocal calls
        calls += 1
        return old_version if calls <= 2 else new_version

    monkeypatch.setattr(api, "latest_committed_version", version_at_read)
    payload = TestClient(api.app).get("/api/projects/Alpha/workspace", params={"view": "overview"}).json()
    assert calls == 3
    assert payload["dataVersion"]["committedImportRef"] == new_version.import_ref
    matching = [key for key in api._WORKSPACE_CACHE if key[0] == str(db.resolve())]
    assert {key[2] for key in matching} == {new_version.run_id}

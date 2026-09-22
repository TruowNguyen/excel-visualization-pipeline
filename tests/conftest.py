from __future__ import annotations

from pathlib import Path
import shutil
from uuid import uuid4

import pytest
from openpyxl import Workbook


@pytest.fixture
def storage_workspace() -> Path:
    root = Path(__file__).resolve().parent / "_storage_runtime"
    path = root / uuid4().hex
    path.mkdir(parents=True)
    try:
        yield path
    finally:
        shutil.rmtree(path, ignore_errors=True)
        try:
            root.rmdir()
        except OSError:
            pass


@pytest.fixture
def sample_workbook() -> Path:
    # Keep the fixture inside the repository so tests also run in restricted sandboxes.
    path = Path(__file__).resolve().parent / "_generated_sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.merge_cells("A3:A4")
    sheet.merge_cells("B3:B4")
    sheet.merge_cells("C3:C4")
    sheet.merge_cells("D3:F3")
    sheet.merge_cells("G3:I3")
    sheet["A3"] = "STT"
    sheet["B3"] = "Dự án"
    sheet["C3"] = "Đơn vị"
    sheet["D3"] = "Kết quả triển khai 12/09/2026"
    sheet["G3"] = "Kết quả triển khai 13/09/2026"
    for col, metric in zip("DEFGHI", ["Tổng số", "Báo sai/Lỗi", "% báo sai"] * 2):
        sheet[f"{col}4"] = metric
    sheet.append([1, "Alpha", "Cảnh báo"])
    sheet.append([None, "1.1. Chất lượng cảnh báo"])
    sheet.append([None, "- Camera", None, 100, 8, 0.08, 120, 6, 0.05])
    sheet["F7"].number_format = "0.00%"
    sheet["I7"].number_format = "0.00%"
    workbook.save(path)
    try:
        yield path
    finally:
        path.unlink(missing_ok=True)
